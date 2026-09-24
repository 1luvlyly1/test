# Databricks notebook source
# MAGIC %pip install openpyxl pdfplumber
# dbutils.library.restartPython()

# COMMAND ----------

CONFIG = {
    "endpoint_name": "databricks-claude-sonnet-4-6",
    "max_tokens": 4096,
    "temperature": 0.0,
    "enterprise_table": "main.default.enterprise_info",
    "dkkd_folder": "/Volumes/main/default/reports/dkkd",
    "report_pdf_path": "/Volumes/main/default/reports/bao_cao_cu.pdf",
    "excel_path": "/Volumes/main/default/reports/Book1.xlsx",
    "schema_scan_rows": 25,
    "schema_cache_path": "/Volumes/main/default/reports/_schema_cache.json",
    "use_schema_cache": True,
    "out_main_md": "/Volumes/main/default/reports/bao_cao_tong_hop.md",
    "out_excel_md": "/Volumes/main/default/reports/du_lieu_excel.md",
    "out_summary_md": "/Volumes/main/default/reports/tong_ket_thay_doi.md",
    "portrait_fields": ["ngành nghề", "hoạt động kinh doanh", "quy trình sản xuất"],
}

# COMMAND ----------

import re, os, json, unicodedata
import openpyxl, pdfplumber
import mlflow
from datetime import datetime
from mlflow.deployments import get_deploy_client

_client = get_deploy_client("databricks")


def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", str(s))
                   if unicodedata.category(c) != "Mn").lower().strip()


def to_local(path):
    p = path[len("dbfs:"):] if path.startswith("dbfs:") else path
    if p.startswith("/Volumes/") or p.startswith("/dbfs/"):
        return p
    return "/dbfs" + p if p.startswith("/") else p


def call_ai(system, user):
    resp = _client.predict(
        endpoint=CONFIG["endpoint_name"],
        inputs={
            "messages": [
                {"role": "system", "content": system},
                {"role": "user", "content": user},
            ],
            "max_tokens": CONFIG["max_tokens"],
            "temperature": CONFIG["temperature"],
        },
    )
    return resp["choices"][0]["message"]["content"]


def parse_json(out):
    out = re.sub(r"^```(?:json)?|```$", "", out.strip(), flags=re.MULTILINE).strip()
    return json.loads(out)


def read_pdf(path):
    with pdfplumber.open(to_local(path)) as pdf:
        return "\n".join(p.extract_text() or "" for p in pdf.pages)

# COMMAND ----------

# MAGIC %sql
# MAGIC CREATE TABLE IF NOT EXISTS main.default.enterprise_info (
# MAGIC   tax_code STRING, representative_name STRING, updated_date STRING
# MAGIC );

# COMMAND ----------

# MAGIC %sql
# MAGIC INSERT INTO main.default.enterprise_info VALUES
# MAGIC   ('0101234567', 'Nguyễn Văn A', '202511'),
# MAGIC   ('0101234567', 'Nguyễn Văn A', '202601'),
# MAGIC   ('0101234567', 'Trần Thị B',   '202603');

# COMMAND ----------

def get_enterprise_rows():
    df = spark.sql(f"SELECT tax_code, representative_name, updated_date "
                   f"FROM {CONFIG['enterprise_table']} ORDER BY updated_date")
    return [r.asDict() for r in df.collect()]

# COMMAND ----------

def to_yyyymm(s):
    d = re.sub(r"\D", "", str(s))
    return d[:6] if len(d) >= 6 else d


def list_dkkd_files(folder):
    local = to_local(folder)
    out = []
    for name in os.listdir(local):
        if not name.lower().endswith(".pdf"):
            continue
        m = re.match(r"^(\d{6})_", name)
        if m:
            out.append({"period": m.group(1), "name": name,
                        "path": os.path.join(local, name)})
    return sorted(out, key=lambda x: x["period"])


def extract_legal_reps(text):
    system = "Bóc tách giấy đăng ký kinh doanh tiếng Việt. Chỉ trả JSON, không giải thích."
    user = ('Liệt kê tên người đại diện pháp luật trong nội dung dưới đây. '
            'Chỉ lấy tên có thật trong văn bản, không suy diễn.\n'
            'Trả JSON: {"legal_representatives": ["..."]}\n\n' + text)
    try:
        return parse_json(call_ai(system, user)).get("legal_representatives", [])
    except Exception:
        return []


def run_dkkd_flow():
    rows = get_enterprise_rows()
    files = list_dkkd_files(CONFIG["dkkd_folder"])
    file_reps = [{**f, "legal_representatives": extract_legal_reps(read_pdf(f["path"]))}
                 for f in files]

    matches = []
    for fr in file_reps:
        table_names = [r["representative_name"] for r in rows
                       if to_yyyymm(r["updated_date"]) == fr["period"]]
        matched = any(strip_accents(t) == strip_accents(f)
                      for t in table_names for f in fr["legal_representatives"])
        matches.append({
            "period": fr["period"], "file": fr["name"],
            "ten_file": fr["legal_representatives"], "ten_bang": table_names,
            "trung_khop": matched if table_names else None,
        })

    max_file = max((f["period"] for f in file_reps), default=None)
    max_tbl = max((to_yyyymm(r["updated_date"]) for r in rows), default=None)
    return {
        "so_dong_bang": len(rows), "so_file": len(files),
        "doi_chieu": matches,
        "ngay_lon_nhat": {
            "max_file": max_file, "max_bang": max_tbl,
            "trung_khop": (max_file == max_tbl) if max_file and max_tbl else None,
        },
    }

# COMMAND ----------

def slice_phuong_an(text):
    start = re.search(r"ng[àa]nh\s*ngh[eề]", text, re.IGNORECASE)
    end = re.search(r"quy\s*tr[ìi]nh\s*s[aả]n\s*xu[aấ]t.*", text, re.IGNORECASE)
    if start and end:
        return text[start.start():end.end()].strip()
    if start:
        return text[start.start():].strip()
    return text.strip()


def extract_portrait(text):
    fields = CONFIG["portrait_fields"]
    keys = ", ".join('"%s": "..."' % f for f in fields)
    system = "Bóc tách báo cáo doanh nghiệp tiếng Việt. Chỉ trả JSON, không markdown."
    user = (f"Trích NGUYÊN VĂN nội dung đánh giá cho từng mục: {fields}. "
            'Mục nào không có trong văn bản để giá trị "". Tuyệt đối không suy diễn.\n'
            f"Trả JSON: {{{keys}}}\n\n{text}")
    try:
        return parse_json(call_ai(system, user))
    except Exception:
        return {f: "" for f in fields}

# COMMAND ----------

def grid_text(ws, max_rows):
    lines = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i >= max_rows:
            break
        cells = [f"C{j}={v!r}" for j, v in enumerate(row) if v is not None]
        lines.append(f"R{i}: " + " | ".join(cells))
    return "\n".join(lines)


def detect_schema(ws, sheet_name):
    system = ("Phân tích layout sheet Excel tiếng Việt có header nhiều tầng và dòng rác. "
              "Chỉ trả JSON index, không giải thích, không trả giá trị dữ liệu.")
    user = (f"Sheet '{sheet_name}'. Lưới (Rn=dòng, Cn=cột, index từ 0):\n"
            f"{grid_text(ws, CONFIG['schema_scan_rows'])}\n\n"
            "Xác định cấu trúc bảng chính, bỏ dòng/cột rác. Trả JSON:\n"
            '{"data_start_row": <int>, "columns": {"ten": <int>, "ma": <int hoặc null>, '
            '"groups": {"<ten_nhom>": {"<con>": <int>}}}}\n'
            "Tên nhóm/con: thường, không dấu, gạch dưới "
            "(dau_ky, cuoi_ky, phat_sinh, nhap, xuat, ton; no, co, thanh_tien, so_luong).")
    return parse_json(call_ai(system, user))


def validate_schema(ws, schema):
    cols = schema["columns"]
    maxc, maxr = ws.max_column, ws.max_row
    assert 0 <= schema["data_start_row"] < maxr
    assert isinstance(cols.get("ten"), int) and 0 <= cols["ten"] < maxc
    if cols.get("ma") is not None:
        assert 0 <= cols["ma"] < maxc
    assert cols.get("groups")
    for subs in cols["groups"].values():
        for idx in subs.values():
            assert idx is None or (0 <= idx < maxc)


def read_by_schema(ws, schema):
    rows = list(ws.iter_rows(values_only=True))
    cols = schema["columns"]
    recs = []
    for r in rows[schema["data_start_row"]:]:
        if cols["ten"] >= len(r) or r[cols["ten"]] in (None, ""):
            continue
        rec = {"ten": str(r[cols["ten"]]).strip(),
               "ma": r[cols["ma"]] if cols.get("ma") is not None and cols["ma"] < len(r) else None,
               "groups": {}}
        for g, subs in cols["groups"].items():
            rec["groups"][g] = {s: (r[i] if i is not None and i < len(r) else None)
                                for s, i in subs.items()}
        recs.append(rec)
    return recs


def load_cache():
    p = to_local(CONFIG["schema_cache_path"])
    if CONFIG["use_schema_cache"] and os.path.exists(p):
        try:
            with open(p, encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def save_cache(cache):
    if not CONFIG["use_schema_cache"]:
        return
    p = to_local(CONFIG["schema_cache_path"])
    os.makedirs(os.path.dirname(p), exist_ok=True)
    with open(p, "w", encoding="utf-8") as f:
        json.dump(cache, f, ensure_ascii=False, indent=2)


def classify_sheet(name):
    k = strip_accents(name)
    if "ton" in k:
        return "hang_ton"
    if "thu" in k:
        return "phai_thu"
    if "tra" in k:
        return "phai_tra"
    return k.replace(" ", "_")


def load_excel(path):
    wb = openpyxl.load_workbook(to_local(path), data_only=True)
    cache = load_cache()
    data = {}
    for name in wb.sheetnames:
        ws = wb[name]
        if name not in cache:
            schema = detect_schema(ws, name)
            validate_schema(ws, schema)
            cache[name] = schema
        data[classify_sheet(name)] = read_by_schema(ws, cache[name])
    save_cache(cache)
    return data

# COMMAND ----------

def thongke_hientai(excel_json):
    system = ("Chuyên viên phân tích tín dụng. Thống kê tình hình HIỆN TẠI của doanh nghiệp "
              "(hàng tồn kho, phải thu, phải trả) từ dữ liệu cho sẵn. Chỉ dùng con số có trong dữ liệu, "
              "không bịa, không suy diễn số không tồn tại. Viết đoạn văn có số liệu, chưa so sánh.")
    user = ("Dữ liệu hiện tại (JSON từ Excel):\n"
            f"{json.dumps(excel_json, ensure_ascii=False, indent=2)}")
    return call_ai(system, user).strip()


def sosanh_tudo(portrait_cu, thongke_ht, excel_json):
    fields = CONFIG["portrait_fields"]
    system = (
        "Chuyên viên phân tích tín dụng. So sánh QUÁ KHỨ (báo cáo PDF cũ) với HIỆN TẠI (Excel).\n"
        "QUY TẮC BẮT BUỘC chống bịa đặt:\n"
        "- Chỉ dùng dữ kiện có thật trong nguồn được cấp. Không thêm thông tin không có.\n"
        "- 'Quá khứ' chỉ lấy từ nội dung PDF cũ. Nếu mục đó trong PDF trống -> ghi 'Không có trong báo cáo cũ'.\n"
        "- 'Hiện tại' chỉ lấy từ dữ liệu Excel/thống kê. Nếu Excel không có dữ liệu cho mục đó -> "
        "ghi 'Không có dữ liệu hiện tại'.\n"
        "- Dòng 'So sánh' phải nêu rõ một trong các tình huống: "
        "(a) cả hai phía có dữ liệu -> nêu điểm giống/khác; "
        "(b) quá khứ KHÔNG có, hiện tại CÓ -> ghi rõ là THÔNG TIN MỚI PHÁT SINH; "
        "(c) quá khứ có, hiện tại không -> ghi rõ là thông tin cũ, hiện KHÔNG còn dữ liệu.\n"
        "- Dòng 'Kết luận' phải có cho MỌI mục: nhận định ngắn gọn về mục đó (tăng/giảm/ổn định/mới phát sinh/mất dữ liệu).\n"
        "- Không suy đoán nguyên nhân nếu nguồn không nêu."
    )
    frame = "\n".join(f"- {f}" for f in fields)
    user = (
        f"NGUỒN QUÁ KHỨ (PDF cũ, JSON):\n{json.dumps(portrait_cu, ensure_ascii=False, indent=2)}\n\n"
        f"NGUỒN HIỆN TẠI - thống kê:\n{thongke_ht}\n\n"
        f"NGUỒN HIỆN TẠI - chi tiết (JSON Excel):\n{json.dumps(excel_json, ensure_ascii=False, indent=2)}\n\n"
        f"Với TỪNG mục dưới đây, xuất Markdown:\n{frame}\n\n"
        "#### <tên mục>\n"
        "**Quá khứ:** <nguyên văn PDF cũ, hoặc 'Không có trong báo cáo cũ'>\n"
        "**Hiện tại:** <từ Excel, hoặc 'Không có dữ liệu hiện tại'>\n"
        "**So sánh:** <theo tình huống a/b/c ở quy tắc>\n"
        "**Kết luận:** <nhận định ngắn gọn cho mục này>\n\n"
        "Sau khi xong tất cả các mục, thêm phần:\n"
        "#### Nhận xét cuối cùng\n"
        "<tổng hợp toàn cảnh 2-4 câu: xu hướng chung, điểm mới phát sinh, điểm cần lưu ý; chỉ dựa trên dữ kiện đã nêu>"
    )
    return call_ai(system, user).strip()

# COMMAND ----------

def fmt(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return f"{v:,}" if isinstance(v, (int, float)) else str(v)


def write_file(path, text):
    local = to_local(path)
    os.makedirs(os.path.dirname(local), exist_ok=True)
    with open(local, "w", encoding="utf-8") as f:
        f.write(text)


def md_excel(excel_data):
    label = {"hang_ton": "Hàng tồn kho", "phai_thu": "Công nợ phải thu",
             "phai_tra": "Công nợ phải trả"}
    L = ["# Dữ liệu Excel - Tình hình hiện tại", ""]
    for key, records in excel_data.items():
        L += [f"## {label.get(key, key)}", ""]
        if records:
            groups = list(records[0]["groups"].keys())
            subs = {g: list(records[0]["groups"][g].keys()) for g in groups}
            head = ["Tên", "Mã"] + [f"{g}.{s}" for g in groups for s in subs[g]]
            L.append("| " + " | ".join(head) + " |")
            L.append("| " + " | ".join(["---"] * len(head)) + " |")
            for r in records:
                cells = [r["ten"], fmt(r["ma"])] + [fmt(r["groups"][g].get(s))
                                                    for g in groups for s in subs[g]]
                L.append("| " + " | ".join(cells) + " |")
        L.append("")
    return "\n".join(L)


def md_main(dkkd, portrait_cu, thongke_ht, sosanh):
    L = ["# Báo cáo tóm tắt thông tin doanh nghiệp", "",
         f"*{datetime.now():%Y-%m-%d %H:%M} - {CONFIG['endpoint_name']}*", "",
         "## 1. Đối chiếu Đăng ký kinh doanh", "",
         f"- Số dòng trong bảng: {dkkd['so_dong_bang']}",
         f"- Số file ĐKKD: {dkkd['so_file']}", "",
         "### 1.1. Đối chiếu tên đại diện theo cùng thời điểm", "",
         "| Kỳ | File | Tên trong file | Tên trong bảng | Trùng khớp |",
         "| --- | --- | --- | --- | --- |"]
    for m in dkkd["doi_chieu"]:
        tk = "Trùng" if m["trung_khop"] else ("Lệch" if m["trung_khop"] is False else "-")
        L.append(f"| {m['period']} | {m['file']} | {', '.join(m['ten_file']) or '-'} "
                 f"| {', '.join(m['ten_bang']) or '-'} | {tk} |")
    n = dkkd["ngay_lon_nhat"]
    tk = "Khớp" if n["trung_khop"] else ("Lệch" if n["trung_khop"] is False else "-")
    L += ["", "### 1.2. So sánh ngày lớn nhất", "",
          f"- Ngày lớn nhất từ tên file: {n['max_file'] or '-'}",
          f"- Ngày lớn nhất trong bảng: {n['max_bang'] or '-'}",
          f"- Kết quả: {tk}", "",
          "## 2. Báo cáo quá khứ (OCR PDF cũ)", ""]
    for f in CONFIG["portrait_fields"]:
        L += [f"### {f.capitalize()}", "", portrait_cu.get(f) or "_(không có trong báo cáo cũ)_", ""]
    L += ["## 3. Thống kê tình hình hiện tại", "", thongke_ht, "",
          "> Dữ liệu chi tiết xem file du_lieu_excel.md", "",
          "## 4. So sánh quá khứ - hiện tại", "", sosanh, ""]
    return "\n".join(L)

# COMMAND ----------

dkkd = run_dkkd_flow()
portrait_cu = extract_portrait(slice_phuong_an(read_pdf(CONFIG["report_pdf_path"])))
excel_data = load_excel(CONFIG["excel_path"])
thongke_ht = thongke_hientai(excel_data)
sosanh = sosanh_tudo(portrait_cu, thongke_ht, excel_data)

write_file(CONFIG["out_main_md"], md_main(dkkd, portrait_cu, thongke_ht, sosanh))
write_file(CONFIG["out_excel_md"], md_excel(excel_data))

print("Đã xuất:", CONFIG["out_main_md"], "và", CONFIG["out_excel_md"])

# COMMAND ----------

def tong_ket_thay_doi(dkkd, sosanh):
    n = dkkd["ngay_lon_nhat"]
    lech_dai_dien = [m["period"] for m in dkkd["doi_chieu"] if m["trung_khop"] is False]
    muc_can = ["Đầu ra", "Đầu vào", "Hàng tồn kho", "Thay đổi pháp lý"]
    checklist = "\n".join(f"- {m}" for m in muc_can)
    system = ("Chuyên viên phân tích tín dụng. Bạn được cấp phần SO SÁNH QUÁ KHỨ - HIỆN TẠI đã viết sẵn "
              "(mỗi mục có các dòng Quá khứ / Hiện tại / So sánh / Kết luận). "
              "CHỈ tóm tắt cho ĐÚNG các mục trong checklist dưới đây, BỎ QUA mọi mục khác. "
              "Mỗi mục viết 1 gạch đầu dòng bắt đầu bằng tên mục, dựa trên dòng 'So sánh' và 'Kết luận' của mục đó "
              "(nêu rõ thay đổi và kết luận; nếu là thông tin mới phát sinh thì nói rõ). "
              "Chỉ dùng thông tin CÓ THẬT trong phần được cấp; nếu một mục không xuất hiện, "
              "ghi '<tên mục>: không có thông tin'. TUYỆT ĐỐI không bịa, không thêm mục ngoài checklist. "
              "Sau các gạch đầu dòng, thêm 1 dòng 'Nhận xét chung:' tổng hợp toàn cảnh 4 mục trên trong 1-2 câu.")
    user = ("CHECKLIST CÁC MỤC CẦN TÓM TẮT (chỉ đúng các mục này):\n"
            f"{checklist}\n\n"
            "PHẦN SO SÁNH QUÁ KHỨ - HIỆN TẠI (nguồn duy nhất):\n"
            f"{sosanh}\n\n"
            "Tóm tắt cho từng mục trong checklist (mỗi mục 1 gạch đầu dòng gồm thay đổi + kết luận), "
            "rồi kết bằng dòng 'Nhận xét chung:'.")
    body = call_ai(system, user).strip()

    dkkd_line = (f"- Ngày cập nhật mới nhất: file {n['max_file']} / bảng {n['max_bang']} "
                 f"({'khớp' if n['trung_khop'] else 'lệch' if n['trung_khop'] is False else 'thiếu dữ liệu'})\n"
                 f"- Kỳ lệch tên đại diện: {', '.join(lech_dai_dien) if lech_dai_dien else 'không có'}")
    return body, dkkd_line


body, dkkd_line = tong_ket_thay_doi(dkkd, sosanh)
md = (f"# Tổng kết thay đổi\n\n*{datetime.now():%Y-%m-%d %H:%M}*\n\n"
      f"## Đăng ký kinh doanh\n\n{dkkd_line}\n\n"
      f"## Thay đổi tình hình quan hệ\n\n{body}\n")
write_file(CONFIG["out_summary_md"], md)

print("Đã xuất tổng kết:", CONFIG["out_summary_md"])
print("=" * 50)
print(body)
